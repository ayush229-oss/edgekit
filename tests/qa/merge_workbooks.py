"""Merge the strategy-sweep workbook and the trend_filter investigation
workbook into a single consolidated .xlsx."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from copy import copy

SWEEP_SRC = r"C:\Users\Ayush\projects\edgekit\tests\qa\EdgeKit_Strategy_Sweep.xlsx"
FILTER_SRC = r"C:\Users\Ayush\projects\edgekit\tests\qa\EdgeKit_trend_filter_investigation.xlsx"
OUT = r"C:\Users\Ayush\projects\edgekit\tests\qa\EdgeKit_Strategy_Report.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)


def copy_sheet(src_wb, src_name, dst_wb, dst_name):
    src = src_wb[src_name]
    dst = dst_wb.create_sheet(dst_name)
    for row in src.iter_rows():
        for cell in row:
            new_cell = dst.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                new_cell.font = copy(cell.font)
                new_cell.fill = copy(cell.fill)
                new_cell.alignment = copy(cell.alignment)
                new_cell.border = copy(cell.border)
                new_cell.number_format = cell.number_format
    for col_letter, dim in src.column_dimensions.items():
        dst.column_dimensions[col_letter].width = dim.width
    dst.freeze_panes = "A2" if src.max_row > 1 else None
    return dst


def main():
    sweep_wb = openpyxl.load_workbook(SWEEP_SRC)
    filter_wb = openpyxl.load_workbook(FILTER_SRC)

    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)  # drop the default blank sheet

    # ── Overview sheet ────────────────────────────────────────────────────
    ov = out_wb.create_sheet("OVERVIEW", 0)
    ov["A1"] = "EdgeKit — Strategy Parameter Sweep & Findings"
    ov["A1"].font = TITLE_FONT
    lines = [
        "",
        "Data source: real MT5 XAUUSD M15 (5,000 bars) unless noted otherwise.",
        "",
        "PART 1 — PARAMETER SWEEP (one sheet per template)",
        "  Every slider combination at a 3-point grid (min / default / max) per",
        "  template, backtested against real data. 3,375 total combinations.",
        "  Sheets: ob_fvg_liq, ema_cross, rsi_mr, bb_bounce, donchian, orb,",
        "          macd_cross, vwap_pullback, supertrend, liq_engulf",
        "  SWEEP_SUMMARY — combination counts and runtime per template.",
        "",
        "PART 2 — TREND_FILTER INVESTIGATION (ema_cross 'trend_filter' slider)",
        "  While reviewing the ema_cross sheet, every parameter combo showed the",
        "  trend_filter toggle (True/False) producing IDENTICAL results. Investigated",
        "  whether this holds on choppier data — it does, on everything tested,",
        "  because it's a mathematical property of EMA recursion (proof included),",
        "  not a data coincidence.",
        "  Sheets: FILTER_FINDINGS (proof + explanation)",
        "          FILTER_REAL_DATA_SCAN (27 real-market combos, 10,000+ crosses)",
        "          FILTER_SYNTHETIC_STRESS_TEST (14,404 adversarial whipsaw crosses)",
        "",
        "BOTTOM LINE",
        "  The 'Require trend confirmation' toggle on the EMA Crossover template",
        "  currently has zero effect on any backtest result. See FILTER_FINDINGS",
        "  for the proof and recommended fix options.",
    ]
    for i, text in enumerate(lines, start=2):
        cell = ov.cell(row=i, column=1, value=text)
        if text.isupper() and text.strip():
            cell.font = BOLD
    ov.column_dimensions["A"].width = 95

    # ── Part 1: sweep sheets in template order, then its summary ─────────
    sweep_order = [s for s in sweep_wb.sheetnames if s != "SUMMARY"] + ["SUMMARY"]
    for name in sweep_order:
        dst_name = "SWEEP_SUMMARY" if name == "SUMMARY" else name
        copy_sheet(sweep_wb, name, out_wb, dst_name)

    # ── Part 2: filter investigation sheets ───────────────────────────────
    rename = {
        "SUMMARY": "FILTER_FINDINGS",
        "REAL_DATA_SCAN": "FILTER_REAL_DATA_SCAN",
        "SYNTHETIC_STRESS_TEST": "FILTER_SYNTHETIC_STRESS_TEST",
    }
    for name in filter_wb.sheetnames:
        copy_sheet(filter_wb, name, out_wb, rename[name])

    out_wb.save(OUT)
    print(f"Saved: {OUT}")
    print("Sheets:", out_wb.sheetnames)


if __name__ == "__main__":
    main()
